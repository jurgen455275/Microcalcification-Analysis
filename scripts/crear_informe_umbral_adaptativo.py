from pathlib import Path
import csv

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


BASE = Path(__file__).resolve().parents[1]
OUT = BASE / "results"
WORK = BASE / "build" / "informe_umbral_adaptativo"
WORK.mkdir(parents=True, exist_ok=True)
DOCX = BASE / "docs" / "informe_tecnico_umbral_adaptativo_inbreast.docx"
DOCX.parent.mkdir(parents=True, exist_ok=True)
CHART = WORK / "comparacion_metricas.png"

NAVY = "17365D"
BLUE = "2F75B5"
PALE = "EAF2F8"
GRID = "D9E1F2"
LIGHT = "F7F9FC"
TEXT = RGBColor(0, 0, 0)


def read_metrics():
    rows = []
    with (OUT / "comparacion_tres_algoritmos.csv").open(encoding="utf-8-sig", newline="") as f:
        for row in csv.DictReader(f):
            for key in row:
                if key != "algoritmo":
                    row[key] = float(row[key])
            rows.append(row)
    return {r["algoritmo"]: r for r in rows}


METRICS = read_metrics()


def pct(value):
    return f"{100 * value:.2f} %"


def set_run_font(run, name="Aptos", size=None, bold=None, color=None):
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=90, bottom=90, end=90):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), "4")
        tag.set(qn("w:color"), GRID)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_table(doc, headers, rows, widths=None, font_size=8.6, alignments=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    repeat_header(table.rows[0])
    for i, text in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = str(text)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_shading(cell, NAVY)
        set_cell_margins(cell)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in paragraph.runs:
            set_run_font(run, size=font_size, bold=True, color=RGBColor(255, 255, 255))
    for row_index, values in enumerate(rows):
        row = table.add_row()
        prevent_row_split(row)
        for i, value in enumerate(values):
            cell = row.cells[i]
            cell.text = str(value)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)
            if row_index % 2 == 1:
                set_cell_shading(cell, PALE)
            paragraph = cell.paragraphs[0]
            paragraph.alignment = (
                alignments[i] if alignments else
                (WD_ALIGN_PARAGRAPH.LEFT if i == 0 else WD_ALIGN_PARAGRAPH.CENTER)
            )
            for run in paragraph.runs:
                set_run_font(run, size=font_size, color=TEXT)
    if widths:
        for row in table.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Inches(width)
    set_table_borders(table)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(0)
    return table


def add_heading(doc, text, level=1):
    paragraph = doc.add_heading(text, level=level)
    paragraph.paragraph_format.keep_with_next = True
    return paragraph


def add_caption(doc, text):
    paragraph = doc.add_paragraph(text)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(8)
    for run in paragraph.runs:
        set_run_font(run, size=8.5)
        run.font.italic = True


def add_math(doc, expression):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(6)
    omath_para = OxmlElement("m:oMathPara")
    omath = OxmlElement("m:oMath")
    math_run = OxmlElement("m:r")
    math_text = OxmlElement("m:t")
    math_text.text = expression
    math_run.append(math_text)
    omath.append(math_run)
    omath_para.append(omath)
    paragraph._p.append(omath_para)
    return paragraph


def add_bullet(doc, text, numbered=False):
    style = "List Number" if numbered else "List Bullet"
    paragraph = doc.add_paragraph(text, style=style)
    paragraph.paragraph_format.space_after = Pt(3)
    return paragraph


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)
    set_run_font(run, size=8, color=RGBColor(100, 100, 100))


def load_font(size, bold=False):
    candidates = [
        Path(r"C:\Windows\Fonts\aptos-display.ttf" if bold else r"C:\Windows\Fonts\aptos.ttf"),
        Path(r"C:\Windows\Fonts\arialbd.ttf" if bold else r"C:\Windows\Fonts\arial.ttf"),
    ]
    for candidate in candidates:
        if candidate.exists():
            return ImageFont.truetype(str(candidate), size)
    return ImageFont.load_default()


def create_chart():
    algorithms = [
        ("Original k20", "cpp_original_documento_equivalente_k20"),
        ("C++ actual k70", "cpp_actual_k70"),
        ("Mejorado fijo", "cpp_mejorado_balanced_min3"),
        ("Mejorado adaptativo", "cpp_mejorado_sensitive_min1"),
    ]
    metrics = [
        ("Sensibilidad", "sensitivity_presence", "2F75B5"),
        ("Especificidad", "specificity_presence", "70AD47"),
        ("Precisión", "precision_presence", "ED7D31"),
        ("Localización", "localized_rate", "8064A2"),
    ]
    width, height = 1800, 980
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    title_font = load_font(42, True)
    axis_font = load_font(25)
    label_font = load_font(23)
    small_font = load_font(21)
    draw.text((90, 38), "Comparación de métricas en INbreast", fill="#111111", font=title_font)
    left, top, right, bottom = 170, 145, 1710, 760
    for tick in range(0, 101, 20):
        y = bottom - (bottom - top) * tick / 100
        draw.line((left, y, right, y), fill="#D9E1F2", width=2)
        txt = f"{tick} %"
        bbox = draw.textbbox((0, 0), txt, font=axis_font)
        draw.text((left - 25 - (bbox[2] - bbox[0]), y - 14), txt, fill="#333333", font=axis_font)
    group_width = (right - left) / len(algorithms)
    bar_width = 64
    gap = 12
    for group_index, (label, key) in enumerate(algorithms):
        group_center = left + group_width * (group_index + 0.5)
        total_width = len(metrics) * bar_width + (len(metrics) - 1) * gap
        start_x = group_center - total_width / 2
        for metric_index, (_, metric_key, color) in enumerate(metrics):
            value = METRICS[key][metric_key] * 100
            x0 = start_x + metric_index * (bar_width + gap)
            x1 = x0 + bar_width
            y0 = bottom - (bottom - top) * value / 100
            draw.rectangle((x0, y0, x1, bottom), fill=f"#{color}")
        bbox = draw.textbbox((0, 0), label, font=label_font)
        draw.text((group_center - (bbox[2] - bbox[0]) / 2, bottom + 24), label, fill="#222222", font=label_font)
    draw.line((left, bottom, right, bottom), fill="#555555", width=3)
    legend_y = 860
    x = 210
    for label, _, color in metrics:
        draw.rectangle((x, legend_y, x + 28, legend_y + 28), fill=f"#{color}")
        draw.text((x + 40, legend_y - 2), label, fill="#222222", font=small_font)
        x += 350
    image.save(CHART, dpi=(180, 180))


create_chart()

doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.72)
section.bottom_margin = Inches(0.68)
section.left_margin = Inches(0.72)
section.right_margin = Inches(0.72)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Aptos"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
normal.font.size = Pt(10.3)
normal.font.color.rgb = TEXT
normal.paragraph_format.space_after = Pt(5)
normal.paragraph_format.line_spacing = 1.08
for name, size in (("Title", 23), ("Heading 1", 15), ("Heading 2", 11.5), ("Heading 3", 10.5)):
    style = styles[name]
    style.font.name = "Aptos Display"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = TEXT
    style.paragraph_format.space_before = Pt(10 if name != "Title" else 0)
    style.paragraph_format.space_after = Pt(5)

title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.add_run("Informe técnico del umbral adaptativo para detección de microcalcificaciones en INbreast")
subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle_run = subtitle.add_run("Comparación del algoritmo original y las implementaciones C++")
set_run_font(subtitle_run, size=13, bold=True, color=RGBColor(47, 117, 181))
meta = doc.add_paragraph()
meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
meta.paragraph_format.space_before = Pt(24)
meta_run = meta.add_run("Jurgen Herrera\nUniversidad Nacional de Costa Rica\nUNICIENCIA\n7 de octubre de 2026")
set_run_font(meta_run, size=11)
scope = doc.add_paragraph()
scope.paragraph_format.space_before = Pt(28)
scope.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
scope.add_run(
    "Este informe documenta la verificación del ground truth de INbreast, compara el C++ incluido en el protocolo, "
    "el C++ portátil actual y una implementación C++ experimental, y propone la siguiente etapa de desarrollo. "
    "El análisis mantiene el umbral adaptativo como componente central del proyecto."
)
doc.add_page_break()

add_heading(doc, "Resumen ejecutivo", 1)
p = doc.add_paragraph()
p.add_run("Hallazgo principal. ").bold = True
p.add_run(
    "El conjunto INbreast contiene 410 imágenes. La revisión de las anotaciones identificó 315 imágenes con "
    "microcalcificaciones: 313 tienen regiones individuales y 2 aparecen únicamente como agrupaciones. Entre las "
    "95 imágenes tratadas como negativas por la evaluación existente, 94 tienen clasificación verificable y una, "
    "20588138, no posee archivo XML. Esta última debe registrarse como ground truth no disponible en futuras evaluaciones."
)
p = doc.add_paragraph()
p.add_run("Decisión metodológica. ").bold = True
p.add_run(
    "El modo experimental que alcanzó la mayor precisión de presencia, 82.18 %, utiliza un umbral DoG fijo de 0.10 "
    "y exige al menos tres candidatos. Este resultado demuestra que el preprocesamiento y la regla a nivel de imagen "
    "pueden reducir falsos positivos, pero no valida el objetivo del proyecto porque reemplaza el umbral adaptativo. "
    "Debe conservarse como ablación y referencia, no como versión final."
)
p = doc.add_paragraph()
p.add_run("Recomendación. ").bold = True
p.add_run(
    "La siguiente versión debe combinar la normalización por percentiles, la exclusión estricta del fondo y la "
    "respuesta multiescala con un umbral adaptativo robusto calculado dentro de la máscara mamaria. La calibración de "
    "k y de la regla de agrupamiento debe hacerse con particiones por paciente y un conjunto de prueba independiente."
)

add_table(doc, ["Dato clave", "Resultado", "Interpretación"], [
    ["Imágenes evaluadas", "410", "Cobertura completa del conjunto"],
    ["Positivas para presencia", "315", "313 individuales más 2 solo cluster"],
    ["Negativas operativas", "95", "94 verificadas más 1 XML ausente"],
    ["ROI individuales", "7 143", "Base para localización por lesión"],
    ["Mejor precisión observada", "82.18 %", "Modo fijo experimental"],
    ["Mayor sensibilidad adaptativa", "98.41 %", "Con especificidad de 3.16 %"],
], widths=[2.15, 1.55, 3.15])

add_heading(doc, "Objetivo y alcance", 1)
doc.add_paragraph(
    "El objetivo es determinar qué implementación reconoce mejor la presencia de microcalcificaciones y establecer "
    "cómo mejorarla sin abandonar el umbral adaptativo. La comparación usa la verdad de referencia disponible en "
    "INbreast y separa tres niveles de resultado: presencia por imagen, localización de al menos una lesión y calidad "
    "de los candidatos. El informe no constituye validación clínica ni sustituye una evaluación externa."
)

add_heading(doc, "Verificación del conjunto INbreast", 1)
doc.add_paragraph(
    "La hoja INbreast.csv aporta los identificadores y metadatos de cada mamografía. La presencia de lesiones se "
    "verificó contra los XML de anotaciones, porque la hoja no contiene por sí sola una etiqueta completa de "
    "microcalcificaciones. Esta distinción evita tratar una fila sin descripción de lesión como negativa automática."
)
add_table(doc, ["Categoría de anotación", "Imágenes", "Detalle"], [
    ["Calcification o Calcifications", "313", "Con regiones individuales"],
    ["Solo Cluster", "2", "20587054 y 20587080"],
    ["Negativas verificables", "94", "Sin anotación positiva de calcificación"],
    ["Cluster total", "21", "27 ROI; 19 también tienen ROI individuales"],
    ["XML ausente", "1", "20588138; registrar como GT no disponible"],
], widths=[2.25, 1.0, 3.55])

add_heading(doc, "Implementaciones comparadas", 1)
add_heading(doc, "C++ original del protocolo", 2)
doc.add_paragraph(
    "El anexo del protocolo usa OpenCV, lee una imagen con imread, divide por 65 535, aplica suavizado gaussiano y "
    "gradientes Sobel, calcula el tensor de estructura y fija k en 20. También usa un espaciado de 0.07 mm por píxel "
    "y una máscara externa denominada breast_mask.png. Extrae candidatos con findContours, contourArea y moments. "
    "El programa trabaja con una imagen y no incorpora evaluación masiva ni comparación con el ground truth."
)
add_heading(doc, "C++ portátil actual", 2)
doc.add_paragraph(
    "La versión actual elimina la dependencia de OpenCV y usa C++17. Recibe archivos .inbr, emplea el espaciado "
    "registrado, permite configurar k, implementa sus propios filtros y extrae componentes conectados de ocho "
    "vecinos. Produce CSV por lote y una DLL para la evaluación contra anotaciones. No reproduce exactamente el "
    "anexo: normaliza por el máximo, mide el área por número de píxeles y calcula el centroide como promedio de las "
    "coordenadas. Estas decisiones explican pequeñas diferencias en el número de candidatos."
)
add_heading(doc, "C++ mejorado experimental", 2)
doc.add_paragraph(
    "La nueva implementación conserva intacto el C++ actual y añade normalización por percentiles 0.5 y 99.5, "
    "enmascaramiento explícito del fondo, erosión de 15 píxeles y una respuesta DoG multiescala con sigmas 0.8, 1.6 "
    "y 3.2. Incluye dos modos. Balanced usa un umbral fijo y requiere tres candidatos; sensitive usa un umbral "
    "adaptativo sobre el producto del tensor y DoG y requiere un candidato. Los dos modos sirven para estudiar el "
    "intercambio entre sensibilidad y especificidad."
)

add_table(doc, ["Aspecto", "Original", "C++ actual", "C++ mejorado"], [
    ["Dependencias", "OpenCV", "C++17", "C++17"],
    ["Entrada", "Imagen única", ".inbr por lote", ".inbr por lote"],
    ["Normalización", "División o min max", "División por máximo", "Percentiles 0.5 y 99.5"],
    ["Máscara", "Archivo externo", "Componente Otsu", "Máscara más erosión"],
    ["Respuesta", "Tensor y DoG", "Tensor y DoG", "Tensor y DoG multiescala"],
    ["Umbral", "Adaptativo k 20", "Adaptativo configurable", "Fijo balanced o adaptativo sensitive"],
    ["Área", "contourArea", "Píxeles del componente", "Píxeles del componente"],
    ["Evaluación", "No incluida", "CSV y DLL", "CSV, DLL y dos modos"],
], widths=[1.35, 1.72, 1.75, 1.95], font_size=7.8)

add_heading(doc, "El umbral adaptativo como núcleo del método", 1)
doc.add_paragraph(
    "El principio del proyecto es que el umbral responda al contenido de cada mamografía. La intensidad, densidad "
    "mamaria, adquisición y preprocesamiento cambian entre imágenes; por ello, una constante global puede favorecer "
    "un subconjunto y degradar otro. En el método actual, la respuesta S se resume dentro de la máscara mamaria M y "
    "el umbral se define a partir de su media y desviación estándar."
)
add_math(doc, "T_adapt = min(0.95, μ_M(S) + k σ_M(S))")
doc.add_paragraph(
    "Esta fórmula es adaptativa, pero la media y la desviación estándar reaccionan a valores extremos y a errores de "
    "segmentación. Si el fondo o el borde mamario entran en S, el umbral deja de representar el tejido relevante. "
    "Por eso la mejora prioritaria no es sustituir la adaptación, sino calcularla sobre una señal y una máscara más "
    "limpias."
)
doc.add_paragraph(
    "Una alternativa robusta usa la mediana y la desviación absoluta mediana. El factor 1.4826 permite expresar la "
    "MAD en una escala comparable con la desviación estándar bajo una distribución aproximadamente normal."
)
add_math(doc, "T_robusto = mediana_M(S) + k · 1.4826 · MAD_M(S)")
doc.add_paragraph(
    "Otra opción es seleccionar un cuantil de la respuesta dentro de la máscara. Esta forma controla directamente "
    "la proporción de píxeles que supera el umbral y puede reducir la dependencia de la escala absoluta."
)
add_math(doc, "T_cuantil = Q_q(S | M)")

add_heading(doc, "Resultados comparativos", 1)
doc.add_paragraph(
    "La tabla reproduce la evaluación existente con 315 imágenes positivas y 95 negativas operativas; este último "
    "grupo incluye la imagen sin XML. La precisión mide cuántas predicciones positivas fueron "
    "correctas; la sensibilidad mide cuántas imágenes positivas se detectaron; la especificidad mide cuántas "
    "negativas se rechazaron correctamente. Debido al desbalance del conjunto, la exactitud sola puede ocultar un "
    "mal comportamiento en las imágenes negativas. La exactitud balanceada pondera sensibilidad y especificidad por "
    "igual."
)

result_keys = [
    ("Original equivalente k20", "cpp_original_documento_equivalente_k20"),
    ("Python referencia k70", "python_referencia_k70"),
    ("C++ actual k70", "cpp_actual_k70"),
    ("Mejorado fijo min3", "cpp_mejorado_balanced_min3"),
    ("Mejorado adaptativo min1", "cpp_mejorado_sensitive_min1"),
]
rows = []
for label, key in result_keys:
    r = METRICS[key]
    rows.append([
        label, f"{int(r['tp'])}/{int(r['tn'])}", f"{int(r['fp'])}/{int(r['fn'])}",
        pct(r["precision_presence"]), pct(r["sensitivity_presence"]),
        pct(r["specificity_presence"]), pct(r["balanced_accuracy"]),
    ])
add_table(doc, ["Implementación", "TP/TN", "FP/FN", "Precisión", "Sensibilidad", "Especificidad", "Exactitud bal"],
          rows, widths=[1.7, 0.7, 0.7, 0.85, 0.85, 0.85, 0.9], font_size=7.4)
add_caption(doc, "Tabla 4 Resultados de presencia por imagen sobre 410 mamografías")

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run().add_picture(str(CHART), width=Inches(6.85))
add_caption(doc, "Figura 1 Comparación de sensibilidad, especificidad, precisión y localización")

add_heading(doc, "Lectura de los resultados", 2)
doc.add_paragraph(
    "El original equivalente detecta 312 de 315 imágenes positivas, pero solo rechaza correctamente 1 de las 95 "
    "negativas. Su sensibilidad de 99.05 % convive con una especificidad de 1.05 %, por lo que casi cualquier imagen "
    "se clasifica como positiva. El C++ actual con k igual a 70 reduce candidatos y falsos positivos, pero su "
    "especificidad sigue en 8.42 % y la exactitud balanceada queda en 48.02 %."
)
doc.add_paragraph(
    "El modo mejorado fijo ofrece el mejor equilibrio observado: precisión de 82.18 %, especificidad de 48.42 % y "
    "exactitud balanceada de 60.08 %. A cambio, la sensibilidad baja a 71.75 %. Este resultado muestra el beneficio "
    "del preprocesamiento y de exigir evidencia agrupada, pero no demuestra la superioridad de un umbral adaptativo."
)
doc.add_paragraph(
    "El modo mejorado adaptativo recupera 310 imágenes positivas y alcanza 98.41 % de sensibilidad. Sin embargo, "
    "clasifica como positivas 92 de las 95 imágenes negativas. Su especificidad de 3.16 % confirma que el umbral "
    "adaptativo actual es demasiado permisivo cuando se usa con una regla de presencia de un solo candidato."
)

add_heading(doc, "Localización y calidad de candidatos", 1)
candidate_rows = []
for label, key in result_keys:
    r = METRICS[key]
    candidate_rows.append([
        label,
        f"{int(r['localized_images'])}/313",
        pct(r["localized_rate"]),
        f"{int(r['total_candidates']):,}".replace(",", " "),
        pct(r["candidate_precision"]),
        f"{r['fp_candidates_per_image']:.2f}",
    ])
add_table(doc, ["Implementación", "Imágenes localizadas", "Tasa", "Candidatos", "Precisión candidato", "FP candidato por imagen"],
          candidate_rows, widths=[1.75, 1.05, 0.75, 0.8, 1.05, 1.1], font_size=7.5)
add_caption(doc, "Tabla 5 Resultados de localización y candidatos")
doc.add_paragraph(
    "El modo mejorado adaptativo localiza al menos una región en 130 de las 313 imágenes con ROI individuales, la "
    "mayor cobertura de los métodos comparados. El modo fijo localiza 111, pero obtiene la mejor precisión por "
    "candidato, 33.78 %. Esta diferencia sugiere que la mejora debe separar dos decisiones: primero, qué píxeles "
    "superan el umbral adaptativo; segundo, cuándo un conjunto de candidatos constituye una detección a nivel de imagen."
)

add_heading(doc, "Diagnóstico técnico", 1)
add_heading(doc, "Normalización DICOM", 2)
doc.add_paragraph(
    "Las 410 imágenes revisadas carecen de WindowCenter y WindowWidth, usan MONOCHROME2 y almacenan 16 bits. La "
    "mediana del valor crudo es cero. El respaldo actual usa esa mediana como centro y el rango como ancho; en esas "
    "condiciones, el fondo cero se transforma en 0.5 en lugar de permanecer en cero. Esto introduce respuestas "
    "artificiales en el fondo y en el límite de la mama. La normalización por percentiles dentro de la máscara evita "
    "que unos pocos valores extremos controlen toda la escala."
)
add_heading(doc, "Máscara mamaria", 2)
doc.add_paragraph(
    "El código actual usa la máscara para calcular la media y la desviación del umbral, pero no restringe de forma "
    "consistente la imagen binaria final antes de contar componentes. Además, la función descrita como eliminación "
    "pectoral selecciona el componente Otsu más grande y aplica cierre morfológico; no segmenta realmente el músculo "
    "pectoral. Estas dos condiciones favorecen candidatos de borde, sobre todo en vistas MLO."
)
add_heading(doc, "Equivalencia entre implementaciones", 2)
doc.add_paragraph(
    "En una muestra de cinco imágenes, el C++ portátil produjo un número distinto de candidatos en tres casos. La "
    "diferencia es coherente con el uso de área por píxeles en lugar de contourArea y con cambios en bordes y kernels. "
    "Antes de atribuir una mejora al lenguaje o a la optimización, conviene igualar la definición matemática de cada "
    "etapa y comparar mapas intermedios."
)
add_heading(doc, "Rendimiento", 2)
doc.add_paragraph(
    "El C++ portátil actual promedió aproximadamente 587.9 ms por imagen en la corrida completa. La evaluación dual "
    "del C++ mejorado promedió cerca de 7 331.8 ms por imagen porque calcula varias escalas y dos modos. En el "
    "benchmark inicial de cinco imágenes, el C++ portátil fue entre 1.5 y 1.75 veces más lento que Python con OpenCV. "
    "La prioridad inmediata debe ser validar la precisión del método adaptativo; la optimización debe realizarse "
    "después de fijar la definición del algoritmo."
)

add_heading(doc, "Arquitectura recomendada para la siguiente versión", 1)
doc.add_paragraph(
    "La siguiente versión debe mantener una sola ruta adaptativa y tratar los modos actuales como comparadores. La "
    "secuencia propuesta es la siguiente."
)
flow_rows = [
    ["1", "Lectura DICOM", "Conservar BitsStored, fotometría y espaciado"],
    ["2", "Normalización", "Percentiles dentro de la mama y fondo igual a cero"],
    ["3", "Segmentación", "Máscara mamaria y exclusión pectoral verificable"],
    ["4", "Respuesta multiescala", "Tensor de estructura combinado con DoG"],
    ["5", "Umbral adaptativo", "Media y sigma, MAD o cuantil calculado solo en M"],
    ["6", "Filtrado de componentes", "Área física, elongación, contraste y distancia al borde"],
    ["7", "Agrupamiento", "Número, proximidad y densidad espacial de candidatos"],
    ["8", "Decisión por imagen", "Puntaje de cluster calibrado en validación"],
]
add_table(doc, ["Paso", "Etapa", "Control principal"], flow_rows, widths=[0.55, 1.75, 4.6], font_size=8.4)

add_heading(doc, "Mejoras propuestas al umbral adaptativo", 2)
for item in [
    "Calcular la distribución de S únicamente dentro de una máscara erosionada para excluir el borde mamario.",
    "Comparar la forma actual media más k sigma con mediana más k MAD y con un cuantil adaptativo.",
    "Normalizar cada escala antes de fusionarla para que una sigma no domine por magnitud.",
    "Ajustar k con el conjunto de validación y congelarlo antes de medir el conjunto de prueba.",
    "Separar el umbral de píxel de la regla de imagen. Un píxel candidato no debería convertir automáticamente toda la mamografía en positiva.",
    "Usar la cantidad y proximidad de candidatos para identificar clusters, porque las microcalcificaciones clínicamente relevantes suelen evaluarse como agrupaciones.",
]:
    add_bullet(doc, item)

add_heading(doc, "Diseño experimental recomendado", 1)
doc.add_paragraph(
    "El ajuste debe realizarse por paciente para impedir que imágenes relacionadas aparezcan en entrenamiento y "
    "prueba. Una partición práctica reserva pacientes para desarrollo, validación y prueba. El conjunto de validación "
    "selecciona la familia de umbral, k, filtros de componentes y regla de agrupamiento; el conjunto de prueba se usa "
    "una sola vez para el resultado final."
)
add_table(doc, ["Experimento", "Variable", "Criterio de selección"], [
    ["Ablación de preprocesamiento", "Actual frente a percentiles y máscara erosionada", "Especificidad y FP por imagen"],
    ["Familia de umbral", "Media sigma, MAD y cuantil", "Exactitud balanceada con sensibilidad mínima"],
    ["Escalas", "Una escala frente a combinación multiescala", "FROC por lesión"],
    ["Regla de imagen", "1 candidato, N candidatos y puntaje de cluster", "Precisión y especificidad"],
    ["Generalización", "Particiones por paciente", "Intervalos de confianza en prueba"],
], widths=[1.55, 3.05, 2.3], font_size=8.2)

add_heading(doc, "Métricas que deben reportarse", 2)
for item in [
    "Presencia por imagen con TP, TN, FP, FN, sensibilidad, especificidad, precisión, F1 y exactitud balanceada.",
    "Localización por lesión mediante sensibilidad frente a falsos positivos por imagen, presentada como curva FROC.",
    "Precisión de candidatos y distribución de falsos positivos por imagen, no solo el promedio.",
    "Resultados por densidad ACR, vista CC o MLO y lateralidad, con tamaños de muestra visibles.",
    "Intervalos de confianza por remuestreo a nivel de paciente.",
]:
    add_bullet(doc, item)

add_heading(doc, "Plan de trabajo priorizado", 1)
plan_rows = [
    ["1", "Corregir entrada y máscara", "Fondo cero, percentiles intramamarios, borde y pectoral"],
    ["2", "Reproducir equivalencia", "Comparar mapas y componentes entre Python y C++"],
    ["3", "Estudiar umbral adaptativo", "Media sigma, MAD y cuantil con validación por paciente"],
    ["4", "Añadir agrupamiento", "Puntaje de cluster separado del umbral de píxel"],
    ["5", "Evaluar y congelar", "Prueba independiente y reporte FROC"],
    ["6", "Optimizar C++", "Reutilizar buffers, paralelizar escalas y perfilar"],
]
add_table(doc, ["Prioridad", "Trabajo", "Resultado esperado"], plan_rows, widths=[0.8, 2.1, 3.8], font_size=8.5)

add_heading(doc, "Conclusiones", 1)
doc.add_paragraph(
    "La revisión confirma que 315 de las 410 imágenes presentan microcalcificaciones según las anotaciones usadas. "
    "El C++ original equivalente y el modo adaptativo mejorado alcanzan sensibilidades cercanas a 99 %, pero fallan "
    "en casi todas las imágenes negativas. El modo fijo mejora la precisión y la especificidad, aunque reduce la "
    "sensibilidad y se aparta del fundamento del proyecto."
)
doc.add_paragraph(
    "La evidencia disponible no justifica reemplazar el umbral adaptativo. Justifica rediseñarlo. La ruta recomendada "
    "es conservar la adaptación por imagen, hacerla resistente a valores extremos, calcularla dentro de una máscara "
    "correcta y separar la selección de píxeles de la decisión final por agrupamiento. El código mejorado actual debe "
    "considerarse una plataforma experimental para esa versión, no el algoritmo definitivo."
)

add_heading(doc, "Limitaciones", 1)
for item in [
    "La fila del C++ original se obtuvo con un comportamiento OpenCV equivalente a k igual a 20, porque el anexo no incluye un evaluador masivo listo para el conjunto entregado.",
    "La evaluación de localización usa 313 imágenes con ROI individuales; las dos imágenes solo cluster requieren una regla específica.",
    "La ausencia del XML de 20588138 impide confirmar su clase desde las anotaciones locales.",
    "Los resultados corresponden a este conjunto y a estas reglas de evaluación; no demuestran desempeño clínico ni generalización a otra institución.",
]:
    add_bullet(doc, item)

add_heading(doc, "Archivos y fuentes de reproducibilidad", 1)
add_table(doc, ["Archivo", "Uso"], [
    ["PROTOCOLO-INBREAST-COMPLETO-2026.docx", "Definición y C++ original"],
    ["INbreast.csv y XML", "Identificadores, metadatos y ground truth"],
    ["deform_inbreast.cpp", "C++ portátil actual"],
    ["deform_inbreast_mejorado.cpp", "C++ experimental nuevo"],
    ["comparacion_tres_algoritmos.csv", "Métricas comparativas"],
    ["verificacion_cpp_mejorado_por_imagen.csv", "Resultados detallados por imagen"],
    ["verificacion_microcalcificaciones_inbreast.xlsx", "Auditoría de presencia y anotaciones"],
], widths=[3.25, 3.55], font_size=8.4)

add_heading(doc, "Apéndice de matrices de confusión", 1)
matrix_rows = []
for label, key in result_keys:
    r = METRICS[key]
    matrix_rows.append([label, int(r["tp"]), int(r["tn"]), int(r["fp"]), int(r["fn"]), pct(r["accuracy"]), pct(r["f1_presence"])])
add_table(doc, ["Implementación", "TP", "TN", "FP", "FN", "Exactitud", "F1"], matrix_rows,
          widths=[2.25, 0.55, 0.55, 0.55, 0.55, 1.0, 0.85], font_size=8.0)

footer = section.footer.paragraphs[0]
add_page_number(footer)

doc.core_properties.title = "Informe técnico del umbral adaptativo para detección de microcalcificaciones en INbreast"
doc.core_properties.author = "Jurgen Herrera"
doc.core_properties.subject = "Comparación del C++ original, el C++ actual y el C++ mejorado"
doc.core_properties.keywords = "INbreast, microcalcificaciones, umbral adaptativo, C++, validación"
doc.save(DOCX)
print(DOCX)
