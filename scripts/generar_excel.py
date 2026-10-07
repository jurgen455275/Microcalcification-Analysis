import argparse
import csv
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


REPO_ROOT = Path(__file__).resolve().parents[1]


def read_csv(path):
    with path.open(encoding="utf-8-sig", newline="") as file:
        return list(csv.DictReader(file))


def add_rows_sheet(workbook, title, rows):
    sheet = workbook.create_sheet(title)
    if not rows:
        sheet.append(["Sin datos"])
        return sheet
    headers = list(rows[0].keys())
    sheet.append(headers)
    for row in rows:
        sheet.append([row.get(header, "") for header in headers])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    header_fill = PatternFill("solid", fgColor="17365D")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for column_index, header in enumerate(headers, 1):
        values = [str(header)] + [str(row.get(header, "")) for row in rows[:200]]
        width = min(42, max(10, max(len(value) for value in values) + 2))
        sheet.column_dimensions[get_column_letter(column_index)].width = width
    return sheet


def main():
    parser = argparse.ArgumentParser(description="Genera el libro de verificación de INbreast")
    parser.add_argument("--results", type=Path, default=REPO_ROOT / "results")
    parser.add_argument("--output", type=Path,
                        default=REPO_ROOT / "results" / "verificacion_microcalcificaciones_inbreast.xlsx")
    args = parser.parse_args()

    workbook = Workbook()
    workbook.remove(workbook.active)

    summary_path = args.results / "verificacion_resumen_final.json"
    with summary_path.open(encoding="utf-8") as file:
        summary = json.load(file)
    summary_rows = [{"métrica": key, "valor": value} for key, value in summary.items()
                    if not isinstance(value, (list, dict))]
    add_rows_sheet(workbook, "Resumen", summary_rows)

    sources = [
        ("Por imagen", "verificacion_por_imagen_final.csv"),
        ("Positivas GT", "imagenes_con_microcalcificaciones_gt.csv"),
        ("Por cantidad", "acierto_por_cantidad.csv"),
        ("Comparación", "comparacion_tres_algoritmos.csv"),
        ("C++ mejorado", "verificacion_cpp_mejorado_por_imagen.csv"),
    ]
    for title, filename in sources:
        add_rows_sheet(workbook, title, read_csv(args.results / filename))

    args.output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(args.output)
    print(args.output)


if __name__ == "__main__":
    main()
